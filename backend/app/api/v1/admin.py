import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import List, Dict, Any, Optional
from fastapi import APIRouter, Depends, HTTPException, status, Request, Response

from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func
from sqlalchemy.orm import selectinload

from app.core.database import get_db
from app.core.security import require_portal_admin, get_password_hash
from app.core.audit import log_audit_event
from app.models.models import (
    PortalAdmin, TestCreator, TestAttender, CreatorPermission,
    Company, Assessment, AssessmentAttempt, AuditLog,
    PlatformSetting, ExternalPlatform, ProctoringEvent
)
from app.schemas.schemas import (
    CreatorCreateRequest, CreatorUpdateRequest,
    PlatformSettingUpdate, ExternalPlatformUpdate,
    EmailSettingsSchema, TestEmailRequest
)
from app.services.platform_adapters.adapters import PlatformAdapterRegistry
from app.services.email_service import EmailService

router = APIRouter(prefix="/admin", tags=["Portal Admin Operations"])

# ----------------- Analytics & Overview -----------------

@router.get("/overview")
async def get_admin_overview(
    admin: PortalAdmin = Depends(require_portal_admin),
    db: AsyncSession = Depends(get_db)
):
    """Retrieves high-level platform statistics and telemetry."""
    creators_cnt = await db.scalar(select(func.count(TestCreator.id)))
    attenders_cnt = await db.scalar(select(func.count(TestAttender.id)))
    assessments_cnt = await db.scalar(select(func.count(Assessment.id)))
    attempts_cnt = await db.scalar(select(func.count(AssessmentAttempt.id)))
    violations_cnt = await db.scalar(select(func.count(ProctoringEvent.id)))
    active_attempts_cnt = await db.scalar(
        select(func.count(AssessmentAttempt.id)).where(AssessmentAttempt.status == "ACTIVE")
    )

    return {
        "total_test_creators": creators_cnt or 0,
        "total_test_attenders": attenders_cnt or 0,
        "total_assessments": assessments_cnt or 0,
        "total_attempts": attempts_cnt or 0,
        "active_attempts_now": active_attempts_cnt or 0,
        "total_proctoring_incidents": violations_cnt or 0,
        "system_status": "OPERATIONAL"
    }

# ----------------- Test Creator Management -----------------

@router.get("/creators")
async def list_creators(
    admin: PortalAdmin = Depends(require_portal_admin),
    db: AsyncSession = Depends(get_db)
):
    """Lists all Test Creators with company info and permissions."""
    res = await db.execute(select(TestCreator))
    creators = res.scalars().all()
    output = []
    for c in creators:
        perm_res = await db.execute(
            select(CreatorPermission).where(
                CreatorPermission.creator_id == c.id,
                CreatorPermission.is_granted == True
            )
        )
        perms = [p.permission_key for p in perm_res.scalars().all()]
        output.append({
            "id": c.id,
            "full_name": c.full_name,
            "email": c.email,
            "phone": c.phone,
            "is_active": c.is_active,
            "company_name": c.company.name if c.company else "Default",
            "permissions": perms,
            "created_at": c.created_at
        })
    return output

@router.post("/creators", status_code=status.HTTP_201_CREATED)
async def create_creator(
    req: CreatorCreateRequest,
    request: Request,
    admin: PortalAdmin = Depends(require_portal_admin),
    db: AsyncSession = Depends(get_db)
):
    """Creates a new Test Creator account with assigned granular permissions."""
    email_clean = req.email.lower().strip()
    res = await db.execute(select(TestCreator).where(TestCreator.email == email_clean))
    if res.scalars().first():
        raise HTTPException(status_code=400, detail="Test Creator email already exists")

    # Find or create company
    company_id = None
    if req.company_name:
        c_res = await db.execute(select(Company).where(Company.name == req.company_name.strip()))
        comp = c_res.scalars().first()
        if not comp:
            comp = Company(name=req.company_name.strip())
            db.add(comp)
            await db.flush()
        company_id = comp.id

    creator = TestCreator(
        full_name=req.full_name.strip(),
        email=email_clean,
        hashed_password=get_password_hash(req.password),
        phone=req.phone,
        company_id=company_id,
        is_active=True
    )
    db.add(creator)
    await db.flush()

    # Assign permissions
    for p_key in req.permissions:
        perm = CreatorPermission(
            creator_id=creator.id,
            permission_key=p_key,
            is_granted=True
        )
        db.add(perm)

    await log_audit_event(
        db=db,
        action="CREATOR_CREATED",
        user_id=admin.id,
        user_role="PORTAL_ADMIN",
        target_type="TestCreator",
        target_id=creator.id,
        ip_address=request.client.host if request.client else None,
        details={"creator_email": creator.email, "permissions": req.permissions}
    )
    await db.commit()

    return {"message": "Test Creator created successfully", "creator_id": creator.id}

@router.patch("/creators/{creator_id}")
async def update_creator(
    creator_id: str,
    req: CreatorUpdateRequest,
    request: Request,
    admin: PortalAdmin = Depends(require_portal_admin),
    db: AsyncSession = Depends(get_db)
):
    """Updates status, permissions, or details of a Test Creator."""
    res = await db.execute(select(TestCreator).where(TestCreator.id == creator_id))
    creator = res.scalars().first()
    if not creator:
        raise HTTPException(status_code=404, detail="Test Creator not found")

    if req.full_name is not None:
        creator.full_name = req.full_name
    if req.phone is not None:
        creator.phone = req.phone
    if req.is_active is not None:
        creator.is_active = req.is_active

    if req.permissions is not None:
        # Re-assign permissions
        del_res = await db.execute(select(CreatorPermission).where(CreatorPermission.creator_id == creator_id))
        for old_p in del_res.scalars().all():
            await db.delete(old_p)
        for p_key in req.permissions:
            new_p = CreatorPermission(creator_id=creator_id, permission_key=p_key, is_granted=True)
            db.add(new_p)

    await log_audit_event(
        db=db,
        action="CREATOR_UPDATED",
        user_id=admin.id,
        user_role="PORTAL_ADMIN",
        target_type="TestCreator",
        target_id=creator.id,
        ip_address=request.client.host if request.client else None,
        details={"is_active": req.is_active, "permissions": req.permissions}
    )
    await db.commit()
    return {"message": "Creator updated successfully"}

# ----------------- Global Platform Configuration -----------------

@router.get("/settings")
async def get_platform_settings(
    admin: PortalAdmin = Depends(require_portal_admin),
    db: AsyncSession = Depends(get_db)
):
    """Retrieves all global platform settings."""
    res = await db.execute(select(PlatformSetting))
    settings_list = res.scalars().all()
    return {s.setting_key: s.setting_value_json for s in settings_list}

@router.post("/settings")
async def update_platform_setting(
    req: PlatformSettingUpdate,
    request: Request,
    admin: PortalAdmin = Depends(require_portal_admin),
    db: AsyncSession = Depends(get_db)
):
    """Creates or updates a global platform setting."""
    res = await db.execute(select(PlatformSetting).where(PlatformSetting.setting_key == req.setting_key))
    setting = res.scalars().first()
    if not setting:
        setting = PlatformSetting(
            setting_key=req.setting_key,
            setting_value_json=req.setting_value,
            description=req.description,
            updated_by=admin.id
        )
        db.add(setting)
    else:
        setting.setting_value_json = req.setting_value
        if req.description:
            setting.description = req.description
        setting.updated_by = admin.id

    await log_audit_event(
        db=db,
        action="SETTING_UPDATED",
        user_id=admin.id,
        user_role="PORTAL_ADMIN",
        target_type="PlatformSetting",
        target_id=req.setting_key,
        ip_address=request.client.host if request.client else None,
        details={"key": req.setting_key, "value": req.setting_value}
    )
    await db.commit()
    return {"message": f"Setting '{req.setting_key}' updated successfully"}

# ----------------- Gmail & SMTP Configuration -----------------

@router.get("/email-settings")
async def get_email_settings(
    admin: PortalAdmin = Depends(require_portal_admin),
    db: AsyncSession = Depends(get_db)
):
    """Retrieves current Gmail / SMTP configuration with safe password masking."""
    cfg = await EmailService.get_smtp_config(db)
    pwd = cfg.get("password", "")
    masked_pwd = ("*" * 10) if pwd else ""
    return {
        "host": cfg.get("host", "smtp.gmail.com"),
        "port": cfg.get("port", 587),
        "user": cfg.get("user", ""),
        "password_masked": masked_pwd,
        "is_configured": cfg.get("is_configured", False),
        "from_name": cfg.get("from_name", "EvalAI Assessment Platform"),
        "from_email": cfg.get("from_email", ""),
        "use_tls": cfg.get("use_tls", True),
        "frontend_url": cfg.get("frontend_url", "http://localhost:5173")
    }

@router.post("/email-settings")
async def update_email_settings(
    req: EmailSettingsSchema,
    request: Request,
    admin: PortalAdmin = Depends(require_portal_admin),
    db: AsyncSession = Depends(get_db)
):
    """Persists Gmail / SMTP credentials and settings to database platform storage."""
    # Retrieve existing settings to keep password if omitted
    existing_cfg = await EmailService.get_smtp_config(db)
    password_to_save = req.password if req.password and req.password.strip() and req.password != ("*" * 10) else existing_cfg.get("password", "")

    setting_val = {
        "host": req.host,
        "port": req.port,
        "user": req.user,
        "password": password_to_save,
        "from_name": req.from_name,
        "from_email": req.from_email or req.user,
        "use_tls": req.use_tls,
        "frontend_url": req.frontend_url
    }

    res = await db.execute(select(PlatformSetting).where(PlatformSetting.setting_key == "email_smtp_config"))
    setting = res.scalars().first()
    if not setting:
        setting = PlatformSetting(
            setting_key="email_smtp_config",
            setting_value_json=setting_val,
            description="Gmail & SMTP configuration for system notifications and verification",
            updated_by=admin.id
        )
        db.add(setting)
    else:
        setting.setting_value_json = setting_val
        setting.updated_by = admin.id

    await log_audit_event(
        db=db,
        action="EMAIL_SETTINGS_UPDATED",
        user_id=admin.id,
        user_role="PORTAL_ADMIN",
        target_type="PlatformSetting",
        target_id="email_smtp_config",
        ip_address=request.client.host if request.client else None,
        details={"user": req.user, "host": req.host, "port": req.port}
    )
    await db.commit()

    return {"message": "Gmail & SMTP configuration updated successfully!"}

@router.post("/test-email")
async def send_test_email(
    req: TestEmailRequest,
    admin: PortalAdmin = Depends(require_portal_admin),
    db: AsyncSession = Depends(get_db)
):
    """Sends a test probe email to verify Gmail credentials and SMTP delivery."""
    cfg = await EmailService.get_smtp_config(db)
    subject = req.subject or "🧠 EvalAI — Gmail SMTP Diagnostic Probe"
    html_content = f"""
    <!DOCTYPE html>
    <html>
    <body style="font-family: sans-serif; background-color: #0b0f17; color: #e2e8f0; padding: 24px;">
      <div style="max-width: 500px; margin: 0 auto; background: #111827; border: 1px solid #1f2937; border-radius: 12px; padding: 24px;">
        <h2 style="color: #38bdf8; margin-top: 0;">✅ Gmail SMTP Probe Successful!</h2>
        <p style="font-size: 14px; line-height: 1.6; color: #cbd5e1;">
          Your EvalAI assessment platform has successfully connected to <strong>{cfg['host']}:{cfg['port']}</strong> and authenticated with sender <strong>{cfg['from_email'] or cfg['user']}</strong>.
        </p>
        <div style="background: #1e293b; padding: 12px; border-radius: 8px; font-size: 12px; font-family: monospace; color: #a5f3fc;">
          {req.custom_message or 'All system reminder and OTP verification emails are ready to dispatch.'}
        </div>
        <p style="font-size: 11px; color: #64748b; margin-top: 20px; text-align: center;">
          EvalAI Portal Diagnostics • Delivered to {req.to_email}
        </p>
      </div>
    </body>
    </html>
    """
    plain_content = f"EvalAI Gmail SMTP Probe: Connected to {cfg['host']}:{cfg['port']}. Diagnostic success."

    res = await EmailService.send_raw_email(
        to_email=req.to_email,
        subject=subject,
        html_body=html_content,
        plain_body=plain_content,
        db=db
    )
    return res


# ----------------- External Coding Platforms -----------------

@router.get("/platforms")
async def list_external_platforms(
    admin: PortalAdmin = Depends(require_portal_admin),
    db: AsyncSession = Depends(get_db)
):
    """Lists external coding platform adapters (HackerRank, LeetCode, Codeforces, AtCoder, CodeChef)."""
    res = await db.execute(select(ExternalPlatform))
    platforms = res.scalars().all()
    if not platforms:
        # Seed default platform rows
        defaults = [
            ("Codeforces", "codeforces", True, "https://codeforces.com/api"),
            ("LeetCode", "leetcode", True, "https://leetcode.com/api"),
            ("HackerRank", "hackerrank", True, "https://www.hackerrank.com/api"),
            ("AtCoder", "atcoder", True, "https://atcoder.jp"),
            ("CodeChef", "codechef", True, "https://api.codechef.com")
        ]
        for name, slug, enabled, endpoint in defaults:
            p = ExternalPlatform(name=name, slug=slug, is_enabled=enabled, api_endpoint=endpoint, status="HEALTHY")
            db.add(p)
        await db.commit()
        res = await db.execute(select(ExternalPlatform))
        platforms = res.scalars().all()

    return platforms

@router.patch("/platforms/{slug}")
async def update_external_platform(
    slug: str,
    req: ExternalPlatformUpdate,
    admin: PortalAdmin = Depends(require_portal_admin),
    db: AsyncSession = Depends(get_db)
):
    """Enables/disables an external coding integration or updates its API endpoint/credentials."""
    res = await db.execute(select(ExternalPlatform).where(ExternalPlatform.slug == slug))
    plat = res.scalars().first()
    if not plat:
        raise HTTPException(status_code=404, detail="Platform not found")

    plat.is_enabled = req.is_enabled
    if req.api_endpoint is not None:
        plat.api_endpoint = req.api_endpoint
    if req.api_key is not None:
        plat.api_key = req.api_key
    if req.rate_limit_per_min is not None:
        plat.rate_limit_per_min = req.rate_limit_per_min

    await db.commit()
    return {"message": f"Platform '{plat.name}' updated successfully"}

@router.post("/platforms/{slug}/test-connection")
async def test_platform_connection(
    slug: str,
    admin: PortalAdmin = Depends(require_portal_admin),
    db: AsyncSession = Depends(get_db)
):
    """Runs a live connection probe to the external platform adapter."""
    res = await db.execute(select(ExternalPlatform).where(ExternalPlatform.slug == slug))
    plat = res.scalars().first()
    if not plat:
        raise HTTPException(status_code=404, detail="Platform not found")

    adapter = PlatformAdapterRegistry.get_adapter(
        platform_slug=plat.slug,
        api_endpoint=plat.api_endpoint,
        api_key=plat.api_key
    )
    test_res = await adapter.test_connection()
    plat.status = test_res.get("status", "HEALTHY")
    await db.commit()

    return test_res

# ----------------- Audit Logs -----------------

@router.get("/audit-logs")
async def list_audit_logs(
    limit: int = 100,
    action: Optional[str] = None,
    admin: PortalAdmin = Depends(require_portal_admin),
    db: AsyncSession = Depends(get_db)
):
    """Retrieves immutable system audit trail entries."""
    query = select(AuditLog).order_by(AuditLog.created_at.desc()).limit(limit)
    if action:
        query = query.where(AuditLog.action == action)
    res = await db.execute(query)
    return res.scalars().all()


# ----------------- User Accounts & Excel Export -----------------

@router.get("/users/all")
async def list_all_users(
    admin: PortalAdmin = Depends(require_portal_admin),
    db: AsyncSession = Depends(get_db)
):
    """Lists all registered users categorized into Admins, Teachers (Creators), and Students (Attenders)."""
    # 1. Admins
    admins_res = await db.execute(select(PortalAdmin))
    admins = admins_res.scalars().all()
    admin_list = [
        {
            "id": a.id,
            "full_name": a.full_name,
            "email": a.email,
            "role": "ADMIN",
            "is_active": a.is_active,
            "created_at": a.created_at.isoformat() if a.created_at else None
        }
        for a in admins
    ]

    # 2. Teachers / Creators
    creators_res = await db.execute(select(TestCreator).options(selectinload(TestCreator.company)))
    creators = creators_res.scalars().all()
    creator_list = [
        {
            "id": c.id,
            "full_name": c.full_name,
            "email": c.email,
            "role": "TEACHER",
            "phone": c.phone or "N/A",
            "company_name": c.company.name if c.company else "Default",
            "is_active": c.is_active,
            "created_at": c.created_at.isoformat() if c.created_at else None
        }
        for c in creators
    ]

    # 3. Students / Attenders
    attenders_res = await db.execute(select(TestAttender))
    attenders = attenders_res.scalars().all()
    attender_list = [
        {
            "id": s.id,
            "full_name": s.full_name,
            "email": s.email,
            "role": "STUDENT",
            "phone": s.phone or "N/A",
            "is_verified": s.is_verified,
            "is_active": s.is_active,
            "created_at": s.created_at.isoformat() if s.created_at else None
        }
        for s in attenders
    ]

    return {
        "admins": admin_list,
        "teachers": creator_list,
        "students": attender_list,
        "total_users": len(admin_list) + len(creator_list) + len(attender_list)
    }


@router.delete("/users/{user_type}/{user_id}")
async def delete_user(
    user_type: str,
    user_id: str,
    request: Request,
    admin: PortalAdmin = Depends(require_portal_admin),
    db: AsyncSession = Depends(get_db)
):
    """Removes a user account (admin, teacher/creator, or student/attender)."""
    user_type_norm = user_type.lower().strip()
    
    if user_type_norm in ["admin", "portal_admin"]:
        if user_id == admin.id:
            raise HTTPException(status_code=400, detail="Cannot delete your own active Portal Admin account.")
        res = await db.execute(select(PortalAdmin).where(PortalAdmin.id == user_id))
        target = res.scalars().first()
        if not target:
            raise HTTPException(status_code=404, detail="Admin account not found.")
        target_role = "PORTAL_ADMIN"
        target_email = target.email
        await db.delete(target)

    elif user_type_norm in ["teacher", "creator", "test_creator"]:
        res = await db.execute(select(TestCreator).where(TestCreator.id == user_id))
        target = res.scalars().first()
        if not target:
            raise HTTPException(status_code=404, detail="Teacher/Creator account not found.")
        target_role = "TEST_CREATOR"
        target_email = target.email
        await db.delete(target)

    elif user_type_norm in ["student", "attender", "test_attender"]:
        res = await db.execute(select(TestAttender).where(TestAttender.id == user_id))
        target = res.scalars().first()
        if not target:
            raise HTTPException(status_code=404, detail="Student/Attender account not found.")
        target_role = "TEST_ATTENDER"
        target_email = target.email
        await db.delete(target)

    else:
        raise HTTPException(status_code=400, detail="Invalid user_type. Expected 'admin', 'teacher', or 'student'.")

    await log_audit_event(
        db=db,
        action="USER_REMOVED",
        user_id=admin.id,
        user_role="PORTAL_ADMIN",
        target_type=target_role,
        target_id=user_id,
        ip_address=request.client.host if request.client else None,
        details={"deleted_user_email": target_email, "user_type": user_type_norm}
    )
    await db.commit()
    return {"message": f"User '{target_email}' successfully removed from system."}


@router.get("/users/export-excel")
async def export_users_excel(
    admin: PortalAdmin = Depends(require_portal_admin),
    db: AsyncSession = Depends(get_db)
):
    """Generates an Excel (.xlsx) file containing 3 separate worksheets for Admins, Teachers, and Students."""
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styling definitions
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10, color="0F172A")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # Demo default passwords mapping for reference in export
    KNOWN_DEMO_PASSWORDS = {
        "admin@assessmentportal.io": "Admin@12345 (Default Demo)",
        "recruiter@techcorp.com": "Creator@12345 (Default Demo)",
        "candidate@example.com": "Candidate@12345 (Default Demo)",
    }

    # Helper function to create sheet
    def build_sheet(sheet_title, headers, rows):
        ws = wb.create_sheet(title=sheet_title)
        ws.views.sheetView[0].showGridLines = True

        # Header Row
        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Data Rows
        for row_data in rows:
            ws.append(row_data)
            row_idx = ws.max_row
            for col_num in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_idx, column=col_num)
                cell.font = data_font
                cell.border = thin_border
                if col_num in [1, 4, 6]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    # 1. Admins Sheet
    admins_res = await db.execute(select(PortalAdmin))
    admins = admins_res.scalars().all()
    admin_headers = ["User ID", "Full Name", "Mail ID / Username", "Password Info", "Status", "Joined Date"]
    admin_rows = []
    for a in admins:
        pwd_info = KNOWN_DEMO_PASSWORDS.get(a.email.lower(), "Bcrypt Hash (Encrypted)")
        admin_rows.append([
            a.id,
            a.full_name,
            a.email,
            pwd_info,
            "Active" if a.is_active else "Inactive",
            a.created_at.strftime("%Y-%m-%d %H:%M") if a.created_at else "N/A"
        ])
    build_sheet("Admins", admin_headers, admin_rows)

    # 2. Teachers / Creators Sheet
    creators_res = await db.execute(select(TestCreator).options(selectinload(TestCreator.company)))
    creators = creators_res.scalars().all()
    teacher_headers = ["User ID", "Full Name", "Mail ID / Username", "Password Info", "Phone Number", "Company", "Status", "Joined Date"]
    teacher_rows = []
    for c in creators:
        pwd_info = KNOWN_DEMO_PASSWORDS.get(c.email.lower(), "Bcrypt Hash (Encrypted)")
        teacher_rows.append([
            c.id,
            c.full_name,
            c.email,
            pwd_info,
            c.phone or "N/A",
            c.company.name if c.company else "Default",
            "Active" if c.is_active else "Inactive",
            c.created_at.strftime("%Y-%m-%d %H:%M") if c.created_at else "N/A"
        ])
    build_sheet("Teachers (Creators)", teacher_headers, teacher_rows)

    # 3. Students / Attenders Sheet
    attenders_res = await db.execute(select(TestAttender))
    attenders = attenders_res.scalars().all()
    student_headers = ["User ID", "Full Name", "Mail ID / Username", "Password Info", "Phone Number", "Verification Status", "Status", "Joined Date"]
    student_rows = []
    for s in attenders:
        pwd_info = KNOWN_DEMO_PASSWORDS.get(s.email.lower(), "Bcrypt Hash (Encrypted)")
        student_rows.append([
            s.id,
            s.full_name,
            s.email,
            pwd_info,
            s.phone or "N/A",
            "Verified" if s.is_verified else "Pending Verification",
            "Active" if s.is_active else "Inactive",
            s.created_at.strftime("%Y-%m-%d %H:%M") if s.created_at else "N/A"
        ])
    build_sheet("Students (Attenders)", student_headers, student_rows)

    # Save workbook to BytesIO stream
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    headers = {
        "Content-Disposition": 'attachment; filename="EvalAI_User_Logins_Report.xlsx"',
        "Access-Control-Expose-Headers": "Content-Disposition"
    }

    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers
    )

